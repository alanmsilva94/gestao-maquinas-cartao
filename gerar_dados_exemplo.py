"""Gera planilhas de exemplo (dados 100% fictícios) para importar no sistema.

Uso:
    pip install openpyxl
    python gerar_dados_exemplo.py

Cria em dados_exemplo/:
    Unidades.xlsx                         (aba Planilha1)   base de estabelecimentos
    Maquinas.xlsx                         (aba BD_Maquinas) parque de máquinas (POS)
    Solicitação de Troca de Maquinas.xlsx (aba Planilha2)   solicitações de troca

Os arquivos têm inconsistências propositais (ID N/D, unidade inexistente, unidade
inativa, série duplicada e modelo vazio) para exercitar a aba "Inconsistências".
"""
import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl

random.seed(7)
PASTA = Path(__file__).resolve().parent / "dados_exemplo"
PASTA.mkdir(exist_ok=True)

# UF real, região real e cidades genéricas ("Cidade 01"...).
UFS = [("SP", "SUDESTE"), ("RJ", "SUDESTE"), ("MG", "SUDESTE"), ("ES", "SUDESTE"),
       ("PR", "SUL"), ("SC", "SUL"), ("RS", "SUL"), ("BA", "NORDESTE"), ("PE", "NORDESTE"),
       ("CE", "NORDESTE"), ("GO", "CENTRO OESTE"), ("DF", "CENTRO OESTE"), ("PA", "NORTE"), ("AM", "NORTE")]
MODELOS = ["MODELO A100", "MODELO B200 PRO", "MODELO C300", "MODELO D400 MINI"]
PROBLEMAS = ["Não carrega nem conecta", "Erro de conexão e não processa o pagamento", "Não imprime comprovante",
             "Não lê cartões", "Não realiza transações (Crédito, Débito, Pix)"]
STATUS_SOL = ["Aguardando abertura", "Chamado aberto", "Em andamento", "Máquina enviada", "Resolvido", "Cancelado"]


def cep():
    return f"{random.randint(1000, 99999):05d}-{random.randint(0, 999):03d}"


def fone():
    return f"(11) 9{random.randint(1000, 9999)}-{random.randint(0, 9999):04d}"


# ------------------------------------------------------------ Unidades
N_UNID = 300
unidades = []
for i in range(1, N_UNID + 1):
    uf, regiao = random.choice(UFS)
    cidade = f"Cidade {random.randint(1, 40):02d}"
    ativo = "Inativo" if i % 25 == 0 else "Ativo"
    unidades.append([
        1000 + i, f"Unidade {i:03d}", f"Rua Exemplo, {random.randint(10, 999)}", f"Unidade {i:03d}",
        f"Bairro {random.randint(1, 30):02d}", uf, cep(), cidade, regiao, f"Sub-região {random.randint(1, 5)}",
        "Unidade", ativo, "Analisada", random.randint(20, 900), fone(), f"Responsável {i:03d}", f"Financeiro {i:03d}",
    ])
CAB_UNID = ["ID", "Nome", "Endereco", "N Fantasia", "Bairro", "Estado", "CEP", "Municipio", "Desc.Região",
            "Sub Regiao", "Desc Unidade", "Status", "Status Anali", "Qtd.Membros", "Tel.Atualiz",
            "Nome Responsavel", "Resp. Financ"]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Planilha1"
ws.append(CAB_UNID); [ws.append(r) for r in unidades]
wb.save(PASTA / "Unidades.xlsx")

# ------------------------------------------------------------ Máquinas
maquinas = []
N_MAQ = 360
for i in range(1, N_MAQ + 1):
    uid = random.choice(unidades)[0]
    maquinas.append([5000 + i, f"SN{random.randint(10**9, 10**10 - 1)}", random.choice(MODELOS), uid])
maquinas[3][3] = "N/D"                      # sem unidade
maquinas[7][3] = 99999                      # unidade inexistente
maquinas[11][1] = maquinas[12][1]          # série duplicada
maquinas[15][2] = ""                       # modelo vazio
maquinas[19][3] = [u[0] for u in unidades if u[11] == "Inativo"][0]   # unidade inativa
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BD_Maquinas"
ws.append(["PDV", "Nº de série", "Modelo", "ID_Unidade"]); [ws.append(r) for r in maquinas]
wb.save(PASTA / "Maquinas.xlsx")

# ------------------------------------------------------------ Solicitações
uni_por_id = {u[0]: u for u in unidades}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Planilha2"
ws.append(["Nº", "Nº de Série", "Modelo", "Problema Reportado", "Estado", "Cidade", "Bairro", "Endereço", "CEP",
           "Responsável", "Contato", "Protocolo", "Status"])
amostra = [m for m in maquinas if m[3] in uni_por_id and m[2]][:60:4]
for n, m in enumerate(amostra, start=1):
    u = uni_por_id[m[3]]
    dt = date.today() - timedelta(days=random.randint(1, 90))
    ws.append([m[0], m[1], m[2], random.choice(PROBLEMAS), u[5], u[7], u[4], u[2], u[6], u[15], u[14],
               2026000000 + n, f"{dt:%d/%m/%Y} - {random.choice(STATUS_SOL)}"])
wb.save(PASTA / "Solicitação de Troca de Maquinas.xlsx")

print(f"{len(unidades)} unidades, {len(maquinas)} máquinas, {len(amostra)} solicitações -> {PASTA}")
